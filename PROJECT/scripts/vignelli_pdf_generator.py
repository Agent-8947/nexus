import os
from pathlib import Path
from docx import Document
from fpdf import FPDF
import datetime

# --- VIGNELLI STYLE CONFIG ---
class VignelliPDF(FPDF):
    def __init__(self, lang_code='EN'):
        super().__init__()
        self.lang_code = lang_code
        self.set_margins(25, 20, 25)
        
        # Load System Fonts for Unicode (RU/UA) support
        font_bold = r'C:\Windows\Fonts\arialbd.ttf'
        font_regular = r'C:\Windows\Fonts\arial.ttf'
        
        if os.path.exists(font_bold) and os.path.exists(font_regular):
            self.add_font('Arial', '', font_regular) # Regular
            self.add_font('Arial', 'B', font_bold)  # Bold
            self.use_unicode = True
        else:
            self.use_unicode = False

    def header(self):
        self.set_y(10)
        self.set_font('Arial' if self.use_unicode else 'helvetica', 'B', 8)
        self.set_text_color(0)
        self.cell(0, 5, 'NEXUS COGNITIVE OS // DOCUMENTATION HUB', 0, 1, 'L')
        self.set_draw_color(0)
        self.set_line_width(0.6)
        self.line(25, 17, 185, 17)
        self.ln(15)

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial' if self.use_unicode else 'helvetica', 'B', 8)
        self.set_text_color(160)
        self.cell(0, 10, f'DOCUMENT: {self.lang_code} // PAGE {self.page_no()}', 0, 0, 'R')

    def add_vignelli_title(self, title):
        self.set_font('Arial' if self.use_unicode else 'helvetica', 'B', 48)
        self.set_text_color(0)
        self.multi_cell(0, 16, title.upper(), 0, 'L')
        self.ln(10)

    def add_section_header(self, text):
        self.ln(8)
        self.set_font('Arial' if self.use_unicode else 'helvetica', 'B', 16)
        self.set_text_color(0)
        self.cell(0, 10, text.upper(), 0, 1, 'L')
        self.ln(2)

    def add_body_text(self, text):
        self.set_font('Arial' if self.use_unicode else 'helvetica', '', 10)
        self.set_text_color(0)
        self.multi_cell(0, 6, text, 0, 'L')
        self.ln(2)

def process_language(lang_code):
    in_file = Path(f"INBOX/{lang_code}/NEXUS_{lang_code}_COMBINED.docx")
    out_dir = Path(f"INBOX/{lang_code}")
    out_file = out_dir / f"NEXUS_{lang_code}_VIGNELLI.pdf"
    
    if not in_file.exists():
        print(f"[-] Missing: {in_file}")
        return

    print(f"[*] Generating Vignelli PDF for {lang_code}...")
    try:
        doc = Document(in_file)
        pdf = VignelliPDF(lang_code)
        
        # COVER PAGE
        pdf.add_page()
        pdf.set_y(100)
        pdf.add_vignelli_title(f"NEXUS {lang_code}")
        pdf.set_font('Arial' if pdf.use_unicode else 'helvetica', 'B', 14)
        pdf.cell(0, 10, f"CONSOLIDATED OPERATIONAL FRAMEWORK // MONOCHROME v1.0", 0, 1, 'L')
        
        # CONTENT PAGES
        pdf.add_page()
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            
            if text.startswith('--- Document:'):
                header_text = text.replace('--- Document: ', '').replace(' ---', '')
                pdf.add_section_header(header_text)
            elif para.style.name.startswith('Heading'):
                pdf.add_section_header(text)
            else:
                pdf.add_body_text(text)

        pdf.output(str(out_file))
        print(f"✅ Generated: {out_file}")
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"❌ Error processing {lang_code}: {e}")

if __name__ == "__main__":
    for lang in ['EN', 'RU', 'UA']:
        process_language(lang)
